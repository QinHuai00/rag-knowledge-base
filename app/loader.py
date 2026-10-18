"""文档解析：文件 → [Page]。

为什么要保留 source / page / section：
答案要能溯源。用户点开 [1] 能看到"员工手册.md · 2.3 试用期与转正"，
这比"模型说试用期 3 个月"可信得多。
"""
from __future__ import annotations

import logging
import re
from dataclasses import dataclass
from pathlib import Path

logger = logging.getLogger(__name__)

SUPPORTED_SUFFIXES = {".md", ".txt", ".docx", ".pdf"}


@dataclass
class Page:
    text: str
    source: str
    page: int | None = None      # PDF 页码；Word/Markdown 为 None
    section: str | None = None   # Markdown/Word 的小节标题，citation 里展示


def clean(text: str) -> str:
    """基础清洗：全角空格、多余空格、连续空行。别过度清洗，标题层级要留着。"""
    text = text.replace("\u3000", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _split_markdown_sections(text: str) -> list[tuple[str | None, str]]:
    """按 Markdown 标题切节，返回 [(section_title, section_text)]。"""
    heading = re.compile(r"^(#{1,4})\s+(.+?)\s*$", re.MULTILINE)
    matches = list(heading.finditer(text))
    if not matches:
        return [(None, text)]

    sections: list[tuple[str | None, str]] = []
    preamble = text[: matches[0].start()].strip()
    if preamble:
        sections.append((None, preamble))
    for i, match in enumerate(matches):
        start = match.start()
        end = matches[i + 1].start() if i + 1 < len(matches) else len(text)
        title = match.group(2).strip()
        sections.append((title, text[start:end].strip()))
    return sections


def _load_text_file(path: Path) -> list[Page]:
    raw = clean(path.read_text(encoding="utf-8", errors="ignore"))
    if not raw:
        raise ValueError(f"{path.name} 是空文件")
    return [
        Page(text=section_text, source=path.name, page=None, section=title)
        for title, section_text in _split_markdown_sections(raw)
        if len(section_text) >= 10
    ]


def _load_docx(path: Path) -> list[Page]:
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError("解析 .docx 需要 python-docx：pip install python-docx") from exc

    doc = Document(str(path))
    pages: list[Page] = []
    title: str | None = None
    buffer: list[str] = []

    def flush() -> None:
        nonlocal buffer
        text = clean("\n".join(buffer))
        if text:
            pages.append(Page(text=text, source=path.name, page=None, section=title))
        buffer = []

    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style = (paragraph.style.name or "") if paragraph.style is not None else ""
        is_heading = style.lower().startswith("heading") or style.startswith("标题")
        if is_heading:
            flush()
            title = text
        else:
            buffer.append(text)

    # Word 的表格不在 paragraphs 里，必须单独抽，否则制度类文档会丢关键数据
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                buffer.append(row_text)
    flush()

    if not pages:
        raise ValueError(f"{path.name} 没有可用文本（扫描件或空文档？）")
    return pages


def _load_pdf(path: Path) -> list[Page]:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise RuntimeError("解析 .pdf 需要 pypdf：pip install pypdf") from exc

    reader = PdfReader(str(path))
    pages: list[Page] = []
    for number, pdf_page in enumerate(reader.pages, start=1):
        text = clean(pdf_page.extract_text() or "")
        if len(text) < 50:
            logger.warning("[跳过] %s 第 %d 页没有文本层（可能是扫描件）", path.name, number)
            continue
        pages.append(Page(text=text, source=path.name, page=number))
    if not pages:
        raise ValueError(f"{path.name} 没有可提取的文本，可能是扫描件或加密 PDF")
    return pages


def load(path: str | Path) -> list[Page]:
    """解析单个文件。不支持的格式抛 ValueError，调用方决定提示方式。"""
    file_path = Path(path)
    suffix = file_path.suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise ValueError(f"不支持的文件类型：{suffix}（支持 {sorted(SUPPORTED_SUFFIXES)}）")
    if suffix == ".pdf":
        return _load_pdf(file_path)
    if suffix == ".docx":
        return _load_docx(file_path)
    return _load_text_file(file_path)
